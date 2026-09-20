"""
Convertit un DOCX existant pour le style roman :
- Fusionne les paragraphes consécutifs du corps en un seul paragraphe par bloc,
  en remplaçant les sauts de paragraphe par des retours à la ligne doux (w:br).
- Les blocs sont délimités par : titres de section, sauts de page, page de titre.
- Applique un retrait de première ligne (0,5 cm) sur le style Normal.

Stratégie : modifier le document SOURCE en place (pas de nouveau Document()),
ce qui préserve toutes les relations, styles, et fichiers internes.

Après génération, tu peux ouvrir le DOCX dans Word et remplacer
manuellement les Maj+Entrée par Entrée aux endroits où tu veux un retrait.

Usage :
  python merge_to_softbreaks.py <fichier.docx> [--indent 0.5]
  python merge_to_softbreaks.py <fichier.docx> --dry-run
"""

import argparse
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------------------------------------------------------------------------
# Helpers de détection
# ---------------------------------------------------------------------------

def is_title_para(para):
    """Titre de section : gras ET taille >= 15pt."""
    if not para.text.strip():
        return False
    for run in para.runs:
        if run.bold and run.font.size and run.font.size.pt >= 15:
            return True
    return False


def is_section_header(para):
    """En-têtes centrés spéciaux."""
    text = para.text.strip()
    if not text:
        return False
    specials = ("— La fin première —", "Page(s) manquante(s)", "———")
    if text in specials or text.startswith("Vingt mois après"):
        return True
    if para.alignment == WD_ALIGN_PARAGRAPH.CENTER:
        for run in para.runs:
            if run.bold:
                return True
    return False


def para_is_page_break(para):
    """Ce paragraphe EST-IL un saut de page (w:br type=page) ?"""
    for br in para._element.findall('.//' + qn('w:br')):
        if br.get(qn('w:type')) == 'page':
            return True
    return False


def is_section_boundary(para):
    """Ce paragraphe démarre-t-il un nouveau bloc (titre ou section header) ?"""
    return is_title_para(para) or is_section_header(para)


def title_page_end_index(paragraphs):
    """Index du paragraphe de saut de page qui termine la page de titre."""
    for i, p in enumerate(paragraphs):
        if para_is_page_break(p):
            return i
    return 0


# ---------------------------------------------------------------------------
# Construction d'un w:br doux (retour à la ligne, pas saut de page)
# ---------------------------------------------------------------------------

def make_soft_br():
    """Crée un w:r contenant un w:br sans type → retour à la ligne doux."""
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    r.append(br)
    return r


# ---------------------------------------------------------------------------
# Calcul des blocs
# ---------------------------------------------------------------------------

def compute_blocks(paragraphs, tp_end):
    """
    Retourne une liste de blocs, chaque bloc étant une liste d'indices.
    Bloc d'un seul élément = paragraphe isolé (titre, saut de page, etc.).
    Bloc de plusieurs éléments = paragraphes à fusionner avec w:br.
    """
    blocks = []
    current_block = []

    for i, para in enumerate(paragraphs):

        # Saut de page : bloc isolé, coupe le bloc en cours
        if para_is_page_break(para):
            if current_block:
                blocks.append(current_block)
                current_block = []
            blocks.append([i])
            continue

        # Page de titre : chaque para reste seul
        if i <= tp_end:
            if current_block:
                blocks.append(current_block)
                current_block = []
            blocks.append([i])
            continue

        # Titre ou section header : bloc isolé + coupe
        if is_section_boundary(para):
            if current_block:
                blocks.append(current_block)
                current_block = []
            blocks.append([i])
            continue

        # Paragraphe vide : ignorer
        if not para.text.strip():
            continue

        current_block.append(i)

    if current_block:
        blocks.append(current_block)

    return blocks


# ---------------------------------------------------------------------------
# Fusion en place
# ---------------------------------------------------------------------------

def merge_in_place(doc, indent_cm=0.5):
    """
    Modifie le document en place :
    - Fusionne les blocs de paragraphes en un seul avec w:br entre chaque ligne.
    - Applique le retrait de première ligne au style Normal.
    - Supprime les paragraphes absorbés du body.
    """
    paragraphs = doc.paragraphs
    tp_end = title_page_end_index(paragraphs)
    blocks = compute_blocks(paragraphs, tp_end)

    body = doc.element.body

    # Traiter uniquement les blocs multi-paragraphes
    for block in blocks:
        if len(block) <= 1:
            continue

        first_idx = block[0]
        first_p = paragraphs[first_idx]
        first_elem = first_p._element

        # Retirer le retrait inline éventuel sur le premier paragraphe
        pPr = first_elem.find(qn('w:pPr'))
        if pPr is not None:
            ind = pPr.find(qn('w:ind'))
            if ind is not None:
                pPr.remove(ind)

        # Fusionner les paragraphes suivants dans le premier
        for idx in block[1:]:
            src_para = paragraphs[idx]
            if not src_para.text.strip():
                # Supprimer le paragraphe vide du body
                body.remove(src_para._element)
                continue

            # Ajouter un w:br doux
            first_elem.append(make_soft_br())

            # Copier tous les enfants sauf w:pPr
            for child in list(src_para._element):
                if child.tag == qn('w:pPr'):
                    continue
                first_elem.append(deepcopy(child))

            # Supprimer le paragraphe source du body
            body.remove(src_para._element)

    # Appliquer le retrait de première ligne sur le style Normal
    style = doc.styles['Normal']
    style.paragraph_format.first_line_indent = Cm(indent_cm)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description='Fusionne les paragraphes en retours à la ligne doux pour style roman.'
    )
    parser.add_argument('docx', help='Fichier DOCX à traiter')
    parser.add_argument('--indent', type=float, default=0.5,
                        help='Retrait de première ligne en cm (défaut : 0.5)')
    parser.add_argument('--dry-run', action='store_true',
                        help='Affiche les blocs sans modifier le fichier')
    args = parser.parse_args()

    docx_path = Path(args.docx)
    if not docx_path.exists():
        print(f'Erreur : fichier introuvable : {docx_path}')
        return

    doc = Document(str(docx_path))
    paragraphs = doc.paragraphs
    tp_end = title_page_end_index(paragraphs)

    if args.dry_run:
        blocks = compute_blocks(paragraphs, tp_end)
        n_merged_blocks = 0
        n_merged_paras = 0
        n_isolated = 0

        for block in blocks:
            if len(block) > 1:
                first_text = paragraphs[block[0]].text.strip()
                print(f'  [BLOC {len(block):3d} lignes] {first_text[:65]}...')
                n_merged_blocks += 1
                n_merged_paras += len(block)
            else:
                text = paragraphs[block[0]].text.strip()
                kind = 'saut-page' if para_is_page_break(paragraphs[block[0]]) else 'isole'
                if text:
                    print(f'  [{kind:10s}] {text[:70]}')
                n_isolated += 1

        print(f'\n--- Resultat ---')
        print(f'  Blocs fusionnes    : {n_merged_blocks}  ({n_merged_paras} paras -> {n_merged_blocks} blocs)')
        print(f'  Paragraphes isoles : {n_isolated}')
        print(f'\n(dry-run - fichier non modifie)')
        return

    merge_in_place(doc, indent_cm=args.indent)
    out_path = docx_path.parent / (docx_path.stem + '_roman.docx')
    doc.save(str(out_path))
    print(f'Sauvegarde : {out_path}')
    print(f'Ouvre ce fichier dans Word. Remplace Maj+Entree par Entree')
    print(f'la ou tu veux un nouveau paragraphe avec retrait.')


if __name__ == '__main__':
    main()
